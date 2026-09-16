from django.contrib import admin

# Register your models here.
from django.contrib import admin, messages
from django.urls import path, reverse
from django.shortcuts import render, get_object_or_404, redirect
from django.utils.safestring import mark_safe
from django.utils.html import format_html
from django.db.models import Q
from QuizMaster.middleware import get_current_request

import json
from django import forms
from django_ckeditor_5.widgets import CKEditor5Widget

from .models import (
    VirtualExperiment,
    ApparatusItem,
    TheoryImage,
    ExperimentStep,
    ExperimentQuestion,
    ExpectedMeasurement,
    ExpectedDataAnalysis,
    ExperimentDraft,          # <--- Ensure this is imported
    ExperimentReport,
    ExperimentAnswer,
    ReportEvaluation,
    ExperimentGraph,
    ExpectedExperimentGraph,
    Programme,
    ProgrammeReports,
    ProgrammeExperiment,
    YearReports,
    YEAR_CHOICES,
)


# === Inlines ===

class ApparatusItemInline(admin.TabularInline):
    model = ApparatusItem
    extra = 0
    readonly_fields = ['image_preview']

    def image_preview(self, obj):
        if obj.image:
            return format_html('<img src="{}" width="100" style="border-radius: 4px;" />', obj.image.url)
        return "(No image)"
    image_preview.short_description = "Preview"


class TheoryImageInline(admin.TabularInline):
    model = TheoryImage
    extra = 0
    readonly_fields = ['image_preview']

    def image_preview(self, obj):
        if obj.image:
            return format_html('<img src="{}" width="100" style="border-radius: 4px;" />', obj.image.url)
        return "(No image)"
    image_preview.short_description = "Preview"


class ExperimentStepInline(admin.StackedInline):
    model = ExperimentStep
    extra = 0
    # 🛠️ Removed 'collapse' so Jazzmin tab navigation renders the form fields properly
    classes = ['tab', 'tab-steps']
    readonly_fields = ['image_preview', 'video_preview']
    fields = [
        'step_number', 'instruction',
        'image', 'image_preview',
        'video', 'video_url', 'video_preview'
    ]

    def formfield_for_dbfield(self, db_field, request, **kwargs):
        if db_field.name == 'instruction':
            kwargs['widget'] = CKEditor5Widget(config_name='super_editor')
        return super().formfield_for_dbfield(db_field, request, **kwargs)

    def image_preview(self, obj):
        if obj.image:
            try:
                return format_html(
                    '<img src="{}" width="140" style="border-radius: 6px; border: 1px solid #ccc;" />',
                    obj.image.url
                )
            except Exception:
                return "(Image path error)"
        return "(No image)"
    image_preview.short_description = "Step Image Preview"

    def video_preview(self, obj):
        if obj.video_url:
            url = obj.video_url
            if "youtube.com" in url or "youtu.be" in url:
                embed_url = self._get_youtube_embed_url(url)
                return format_html(
                    '<iframe width="280" height="150" src="{}" '
                    'frameborder="0" '
                    'allow="accelerometer; autoplay; clipboard-write; encrypted-media; gyroscope; picture-in-picture" '
                    'referrerpolicy="strict-origin-when-cross-origin" '
                    'allowfullscreen></iframe>',
                    embed_url
                )
            elif "drive.google.com" in url:
                return format_html(
                    '<iframe src="{}" width="280" height="150" allowfullscreen></iframe>',
                    self._get_drive_embed_url(url)
                )
            else:
                return format_html(
                    '<video width="280" controls><source src="{}" type="video/mp4"></video>',
                    url
                )
        return "(No video)"

    def _get_youtube_embed_url(self, url):
        # Extract video ID cleanly from standard or shortened URLs
        if "watch?v=" in url:
            video_id = url.split("watch?v=")[1].split("&")[0]
        elif "youtu.be/" in url:
            video_id = url.split("youtu.be/")[1].split("?")[0]
        elif "embed/" in url:
            video_id = url.split("embed/")[1].split("?")[0]
        else:
            return url

        # 🛠️ Uses youtube-nocookie.com domain to eliminate Error 153 configuration blocks
        return f"https://www.youtube-nocookie.com/embed/{video_id}"

    def _get_drive_embed_url(self, url):
        if "file/d/" in url and "/view" in url:
            file_id = url.split("/file/d/")[1].split("/")[0]
            return f"https://drive.google.com/file/d/{file_id}/preview"
        return url


class ExperimentQuestionInline(admin.TabularInline):
    model = ExperimentQuestion
    extra = 0
    # 🛠️ Removed 'collapse' here as well
    classes = ['tab', 'tab-questions']
    fields = ['question_text', 'marks']
    
    def formfield_for_dbfield(self, db_field, request, **kwargs):
        if db_field.name == 'question_text':
            kwargs['widget'] = CKEditor5Widget(config_name='super_editor')
        return super().formfield_for_dbfield(db_field, request, **kwargs)


def apparatus_choices_for_experiment(experiment):
    return [("__any__", "Any instrument")] + \
           list(experiment.apparatus_items.values_list("name", "name"))


class ExpectedMeasurementForm(forms.ModelForm):
    class Meta:
        model = ExpectedMeasurement
        fields = "__all__"

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        if self.instance and self.instance.experiment_id:
            choices = apparatus_choices_for_experiment(self.instance.experiment)
        else:
            choices = [("__any__", "Any instrument")]
        self.fields["instrument"].widget = forms.Select(choices=choices)


class ExpectedMeasurementInline(admin.TabularInline):
    model = ExpectedMeasurement
    form = ExpectedMeasurementForm
    extra = 0
    fields = [
        "object_name", "quantity", "instrument",
        "expected_value", "expected_unit", "tolerance",
        "marks", "required", "order",
    ]


class ExpectedDataAnalysisInline(admin.TabularInline):
    model = ExpectedDataAnalysis
    extra = 0
    fields = ["object_name", "quantity", "formula_latex", "alt_formula_latex", "unit", "marks", "order"]
    ordering = ["order"]


# === Forms ===

class VirtualExperimentForm(forms.ModelForm):
    class Meta:
        model = VirtualExperiment
        fields = [
            'title', 'code', 'year_group',
            'objective', 'theory', 'method_summary',
            'recording_guide_tables', 'video_url', 'requires_report', 'requires_questions', 'needs_graph',
            'unlock_time', 'lock_time'
        ]
        widgets = {
            'objective': CKEditor5Widget(config_name='super_editor'),
            'theory': CKEditor5Widget(config_name='super_editor'),
            'method_summary': CKEditor5Widget(config_name='super_editor'),
            'recording_guide_tables': CKEditor5Widget(config_name='super_editor'),
            'unlock_time': forms.SplitDateTimeWidget(
                date_attrs={'type': 'date', 'class': 'vDateField'},
                time_attrs={'type': 'time', 'class': 'vTimeField'}
            ),
        }

@admin.register(VirtualExperiment)
class VirtualExperimentAdmin(admin.ModelAdmin):
    form = VirtualExperimentForm

    list_display = [
        'title', 'code', 'year_group',
        'requires_report', 'requires_questions', 'needs_graph',
        'unlock_time', 'lock_time',
    ]

    list_filter = ['year_group', 'requires_report', 'requires_questions', 'needs_graph', 'unlock_time', 'lock_time']

    fieldsets = (
        ('Basic Information', {
            'fields': (
                'title', 'code', 'year_group',
                'objective', 'theory', 'method_summary',
                'recording_guide_tables',
                'video_url',
            ),
        }),
        ('Scheduling', {
            'fields': ('unlock_time', 'lock_time'),
        }),
        ('Requirements', {
            'fields': ('requires_report', 'requires_questions', 'needs_graph'),
        }),
    )

    inlines = [
        ApparatusItemInline,
        TheoryImageInline,
        ExperimentStepInline,
        ExperimentQuestionInline,
        ExpectedMeasurementInline,
        ExpectedDataAnalysisInline, 
    ]

    def get_queryset(self, request):
        qs = super().get_queryset(request)
        return qs.prefetch_related(
            'apparatus_items',
            'theory_images',
            'experimentstep_set',
            'questions',
            'expected_measurements',
            'expected_data_analysis',
        )

    class Media:
        js = (
            'admin/js/core.js',
            'admin/js/calendar.js',
            'admin/js/admin/DateTimeShortcuts.js',
        )
        css = {
            'all': ( 
                'admin/css/widgets.css',
                'admin/css/admin_custom.css',
            ),
        }



@admin.register(ExperimentDraft)
class ExperimentDraftAdmin(admin.ModelAdmin):
    list_display = ("id", "experiment", "student", "is_active", "last_updated", "short_data")

    def short_data(self, obj):
        if not obj.data:
            return "—"
        # show only first row for preview
        try:
            first_row = obj.data[0]
            return json.dumps(first_row)
        except Exception:
            return str(obj.data)
    short_data.short_description = "Draft Data"




@admin.register(ExperimentGraph)
class ExperimentGraphAdmin(admin.ModelAdmin):
    list_display = (
        "id", "report", "title", "created_at",
        "graph_preview", "points_count", "vector_present"
    )
    readonly_fields = (
        "graph_preview", "points_count", "vector_present", "created_at"
    )
    search_fields = (
        "report__experiment__title",
        "report__student__username",
        "title"
    )
    list_filter = ("created_at",)

    def graph_preview(self, obj):
        if not obj.points_json:
            return "— No data —"

        # Serialize points
        points_json = json.dumps([{'x': p['x'], 'y': p['y']} for p in obj.points_json])
        vector_json = json.dumps(obj.vector_json if isinstance(obj.vector_json, list) else [])

        html = f"""
        <canvas id="graph_{obj.id}" width="220" height="140"></canvas>
        <script>
        (function() {{
            const ctx = document.getElementById('graph_{obj.id}').getContext('2d');
            const points = {points_json};
            const vector = {vector_json};
            const datasets = [{{
                label: 'Data Points',
                data: points,
                backgroundColor: 'rgba(75, 192, 192, 0.6)',
                borderColor: 'rgba(75, 192, 192, 1)',
                showLine: false,
                pointRadius: 4,
            }}];

            if (vector && vector.length > 0) {{
                datasets.push({{
                    label: 'Best Fit Line',
                    data: vector,
                    borderColor: 'rgba(255, 99, 132, 1)',
                    backgroundColor: 'rgba(255, 99, 132, 0.1)',
                    showLine: true,
                    fill: false,
                    pointRadius: 0,
                    borderWidth: 2,
                    tension: 0
                }});
            }}

            new Chart(ctx, {{
                type: 'scatter',
                data: {{ datasets: datasets }},
                options: {{
                    responsive: false,
                    plugins: {{
                        legend: {{ display: false }}
                    }},
                    scales: {{
                        x: {{
                            title: {{ display: true, text: '{obj.x_label or "X"}' }},
                            beginAtZero: true
                        }},
                        y: {{
                            title: {{ display: true, text: '{obj.y_label or "Y"}' }},
                            beginAtZero: true
                        }}
                    }}
                }}
            }});
        }})();
        </script>
        """
        return mark_safe(html)
    graph_preview.short_description = "Graph Preview"

    def points_count(self, obj):
        return len(obj.points_json or [])
    points_count.short_description = "Points"

    def vector_present(self, obj):
        return bool(obj.vector_json)
    vector_present.boolean = True
    vector_present.short_description = "Vector Present"
 
    def has_module_permission(self, request):
            # Keeps the URL active but hides it from the sidebar menu
            return False

class ExperimentAnswerInline(admin.TabularInline):
    model = ExperimentAnswer
    extra = 0
    readonly_fields = ('question', 'answer_text')


class ReportEvaluationInline(admin.StackedInline):
    model = ReportEvaluation
    can_delete = False
    extra = 0
    readonly_fields = ('evaluator', 'evaluated_at', 'display_scores', 'total_score')

    def display_scores(self, obj):
        if not obj or not obj.scores:
            return "No scores available"
        html = "<ul style='list-style:none; padding-left:0; margin:0;'>"
        for section, score in obj.scores.items():
            feedback = obj.feedback.get(section, "No feedback") if obj.feedback else "No feedback"
            html += f"<li><strong>{section.replace('_', ' ').title()}</strong>: {score}/10<br><em>{feedback}</em></li>"
        html += "</ul>"
        return mark_safe(html)
    display_scores.short_description = "Evaluation Scores & Feedback"

    def total_score(self, obj):
        if not obj or not obj.scores:
            return "-"
        total = 0
        for v in obj.scores.values(): 
            if isinstance(v, dict):
                try:
                    total += int(v.get("rubric", 0))
                except (TypeError, ValueError):
                    pass
        return total
    total_score.short_description = "Total Score"


@admin.register(ExperimentReport)
class ExperimentReportAdmin(admin.ModelAdmin):

    list_display = (
        "report_info",
        "score_display",
        "status_display",
        "data_signals",
        "actions_panel",
    )


    list_filter = (
        ('student__programme', admin.RelatedOnlyFieldListFilter), 
        'student__year_group',
        'submitted_at',
    )
    ordering = ['experiment', '-submitted_at']  
    # These fields help the search bar find things across relationships
    search_fields = ("experiment__title", "student__username", "student__programme__name")
    
    inlines = [ExperimentAnswerInline, ReportEvaluationInline]
    actions = ["evaluate_reports", "export_selected_to_excel_zip"]


    # =========================================================================
    # 📥 THE EXPORT ENGINE ACTION
    # =========================================================================

    @admin.action(description="Export selected reports to Excel (ZIP archive)")
    def export_selected_to_excel_zip(self, request, queryset):
        import zipfile
        import re
        import os
        from pathlib import Path
        from io import BytesIO
        from collections import defaultdict
        from tempfile import NamedTemporaryFile
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, GradientFill, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.drawing.image import Image as XLImage
        from django.http import FileResponse
        from django.utils import timezone
        from django.contrib import messages
        from django.conf import settings
        from .models import ReportEvaluation

        report_ids = queryset.values_list("id", flat=True)
        evaluations = ReportEvaluation.objects.filter(report_id__in=report_ids).select_related(
            "report__student", "report__student__programme", "report__experiment", "evaluator"
        )

        if not evaluations.exists():
            self.message_user(request, "None of the selected reports have completed evaluations yet.", level=messages.WARNING)
            return

        grouped_data = defaultdict(lambda: defaultdict(list))
        
        for eval_obj in evaluations:
            exp_title = eval_obj.report.experiment.title
            student_user = eval_obj.report.student
            
            group_key = None
            
            # 🔍 STAGE 1: Check the direct concrete relationship link first (e.g., Bachelor of Science (Physics))
            if student_user.programme and student_user.programme.code:
                group_key = f"Programme Group: {student_user.programme.code.upper()}"
                
            # 🔍 STAGE 2: If the program link is missing, check the Registration Number field column text
            elif student_user.registration_number:
                match = re.search(r"([a-zA-Z]{2,})\s*([0-9]{3})", student_user.registration_number)
                if match:
                    group_key = f"Programme Group: {match.group(1).upper()}{match.group(2)}"
                    
            # 🔍 STAGE 3: If both are missing, parse the Username string itself for a match (e.g., SC209/1234/26)
            if not group_key:
                match = re.search(r"([a-zA-Z]{2,})\s*([0-9]{3})", student_user.username)
                if match:
                    group_key = f"Programme Group: {match.group(1).upper()}{match.group(2)}"
                else:
                    # 🔍 STAGE 4: Final fallback for pure test strings like "Student 001" or "student 05"
                    group_key = "Programme Group: Other"
                
            grouped_data[exp_title][group_key].append(eval_obj)

        # ---- (Keep your standard path resolution and excel style logic below intact) ----
        logo_path = None
        possible_paths = [
            Path(settings.BASE_DIR) / "QuestionBank" / "static" / "images" / "logo.png",
            Path(settings.BASE_DIR) / "static" / "images" / "logo.png",
        ]
        for path in possible_paths:
            if path.exists():
                logo_path = path
                break

        border_std = Border(
            left=Side(style="thin", color="A6A6A6"), right=Side(style="thin", color="A6A6A6"),
            top=Side(style="thin", color="A6A6A6"), bottom=Side(style="thin", color="A6A6A6"),
        )
        fill_header = PatternFill("solid", fgColor="00B4C7E7")
        fill_alt = PatternFill("solid", fgColor="F8FBFD")
        headers = ["Student", "Year Group", "Reg No", "Evaluator", "Submitted", "Evaluated", "Score (/10)"]

        zip_buffer = BytesIO()
        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
            for exp_title, prefixes in sorted(grouped_data.items()):
                wb = Workbook()
                ws = wb.active
                ws.title = "Experiment Reports"
                ws.sheet_view.showGridLines = False

                if logo_path:
                    try:
                        img = XLImage(str(logo_path))
                        img.width, img.height = 100, 100 
                        ws.add_image(img, "A1")
                    except Exception:
                        pass

                ws.merge_cells("B2:H4")
                title_cell = ws.cell(row=2, column=2, value="QUESTION BANK PERFORMANCE REPORT")
                title_cell.font = Font(name="Cambria", size=22, bold=True, color="FFFFFF")
                title_cell.alignment = Alignment(horizontal="center", vertical="center")
                title_cell.fill = GradientFill(type="linear", stop=("1F4E78", "DCE6F1"), degree=90)

                ws.merge_cells("B5:H5")
                subtitle = ws.cell(row=5, column=2, value="Student Achievement Report | Question Bank Analytics")
                subtitle.font = Font(name="Calibri", size=12, italic=True, color="305496")
                subtitle.alignment = Alignment(horizontal="center", vertical="center")

                meta = [
                    f"Report Type: Filtered Admin Selection",
                    f"Experiment: {exp_title}",
                    f"Generated on: {timezone.localtime(timezone.now()).strftime('%Y-%m-%d %H:%M')}",
                ]
                for offset, text in enumerate(meta, 1):
                    ws.cell(row=6 + offset, column=2, value=text).font = Font(size=12, italic=True, color="404040")

                row_num = 10
                ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=len(headers))
                c = ws.cell(row=row_num, column=1, value=f"EXPERIMENT: {exp_title}")
                c.font = Font(name="Cambria", size=14, bold=True, color="FFFFFF")
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.fill = GradientFill(type="linear", degree=0, stop=("1F4E78", "DCE6F1"))
                row_num += 1

                for prefix in sorted(prefixes.keys(), key=lambda x: ("Other" in x, x)):
                    evals = prefixes[prefix]

                    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=len(headers))
                    c = ws.cell(row=row_num, column=1, value=f"{prefix} ({len(evals)} Students)")
                    c.font = Font(name="Calibri", size=12, bold=True, color="1F4E78")
                    c.alignment = Alignment(horizontal="left", indent=1, vertical="center")
                    c.fill = PatternFill("solid", fgColor="DCE6F1")
                    row_num += 1

                    for col_idx, h_text in enumerate(headers, 1):
                        cell = ws.cell(row=row_num, column=col_idx, value=h_text)
                        cell.font = Font(bold=True, color="FFFFFF")
                        cell.fill = fill_header
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                        cell.border = border_std
                        ws.column_dimensions[get_column_letter(col_idx)].width = 22
                    row_num += 1

                    evals.sort(key=lambda e: e.report.student.username)

                    for eval_obj in evals:
                        scores = eval_obj.scores or {}
                        numeric_scores = []
                        for v in scores.values():
                            if isinstance(v, (int, float)): numeric_scores.append(float(v))
                            elif isinstance(v, str) and v.replace(".", "", 1).isdigit(): numeric_scores.append(float(v))
                            elif isinstance(v, dict) and isinstance(v.get("rubric"), (int, float, str)):
                                r_val = str(v.get("rubric"))
                                if r_val.replace(".", "", 1).isdigit(): numeric_scores.append(float(r_val))

                        avg = round(sum(numeric_scores) / len(numeric_scores), 1) if numeric_scores else "-"
                        student_user = eval_obj.report.student

                        row_data = [
                            student_user.get_full_name() or student_user.username,
                            getattr(student_user, "year_group", "—"),
                            student_user.registration_number or student_user.username,  # ✅ Shows actual Reg No if present
                            getattr(eval_obj.evaluator, "username", "—"),
                            timezone.localtime(eval_obj.report.submitted_at).strftime("%Y-%m-%d") if eval_obj.report.submitted_at else "—",
                            timezone.localtime(eval_obj.evaluated_at).strftime("%Y-%m-%d") if eval_obj.evaluated_at else "—",
                            avg,
                        ]

                        for col_idx, val in enumerate(row_data, 1):
                            cell = ws.cell(row=row_num, column=col_idx, value=val)
                            cell.alignment = Alignment(horizontal="center", vertical="center")
                            cell.border = border_std
                            if row_num % 2 == 0: cell.fill = fill_alt
                        row_num += 1
                    row_num += 1

                ws.freeze_panes = "A11"
                
                with NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
                    wb.save(tmp.name)
                    stamp = timezone.localtime(timezone.now()).strftime("%Y%m%d_%H%M%S")
                    clean_title = "".join([c if c.isalnum() else "_" for c in exp_title])
                    zip_file.write(tmp.name, arcname=f"{clean_title}_{stamp}.xlsx")

        zip_buffer.seek(0)
        return FileResponse(
            zip_buffer,
            as_attachment=True,
            filename=f"Admin_Export_{timezone.localtime(timezone.now()).strftime('%Y%m%d_%H%M%S')}.zip",
            content_type="application/zip"
        )


    def get_queryset(self, request):
        return (
            super()
            .get_queryset(request)
            .filter(submitted_at__isnull=False)
      
            .select_related("experiment", "student__programme", "evaluation")
            .defer(
                "group_members", "objective", "theory", "apparatus_scope",
                "procedure", "results", "data_analysis", "discussion",
                "conclusion", "references", "observation", "data",
                "graph_x_values", "graph_y_values",
            )
        )




    def report_info(self, obj):
        return format_html(
            '''
            <div>
                <strong><i class="fas fa-flask"></i> {}</strong><br>
                <span><i class="fas fa-user"></i> {}</span><br>
                <small><i class="fas fa-clock"></i> {}</small>
            </div>
            ''',
            obj.experiment,
            obj.student,
            obj.submitted_at.strftime("%b %d, %Y %H:%M") if obj.submitted_at else "-"
        )
    report_info.short_description = "Report"


    def score_display(self, obj):
        eval_obj = getattr(obj, "evaluation", None)

        if not eval_obj or not getattr(eval_obj, "scores", None):
            return format_html('<span style="color:gray;">—</span>')

        scores = eval_obj.scores or {}
        numeric_scores = [
            float(v.get("rubric")) for v in scores.values()
            if isinstance(v, dict) and v.get("rubric") is not None
        ]

        if not numeric_scores:
            return format_html('<span style="color:gray;">—</span>')

        avg = round(sum(numeric_scores) / len(numeric_scores), 1)

        color = "green" if avg >= 7 else "orange" if avg >= 4 else "red"

        return format_html(
            '<strong style="color:{};"><i class="fas fa-chart-line"></i> {}/10</strong>',
            color, avg
        )
    score_display.short_description = "Score"

    def status_display(self, obj):
        eval_obj = getattr(obj, "evaluation", None)

        if eval_obj:
            return format_html(
                '<span style="color:green;"><i class="fas fa-check-circle"></i> Evaluated</span>'
            )
        else:
            return format_html(
                '<span style="color:orange;"><i class="fas fa-hourglass-half"></i> Pending</span>'
            )
    status_display.short_description = "Status"


    def data_signals(self, obj):
        has_file = bool(obj.report_file)
        file_icon = "fa-file-alt" if has_file else "fa-times-circle"
        file_color = "green" if has_file else "red"
        file_text = "File" if has_file else "No File"

        graph_count = obj.graphs.count()
        graph_color = "green" if graph_count > 0 else "gray"

        return format_html(
            '''
            <div>
                <span style="color:{};">
                    <i class="fas {}"></i> {}
                </span><br>
                <span style="color:{};">
                    <i class="fas fa-chart-bar"></i> {} Graphs
                </span>
            </div>
            ''',
            file_color, file_icon, file_text,
            graph_color, graph_count
        )
    data_signals.short_description = "Data"


    def actions_panel(self, obj):
        preview_url = reverse("admin:experimentreport_preview", args=[obj.id])

        graph_url = reverse("admin:virtuallab_experimentgraph_changelist")
        graph_url += f"?report__id__exact={obj.id}"

        return format_html(
            '''
            <a class="button" href="{}" target="_blank">
                <i class="fas fa-eye"></i>
            </a>
            &nbsp;
            <a class="button" href="{}">
                <i class="fas fa-chart-line"></i>
            </a>
            ''',
            preview_url,
            graph_url
        )
    actions_panel.short_description = "Actions"




    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path(
                '<int:report_id>/preview/',
                self.admin_site.admin_view(self.admin_preview_report),
                name='experimentreport_preview',
            ),
        ]
        return custom_urls + urls

    def preview_report(self, obj):
        url = reverse("admin:experimentreport_preview", args=[obj.id])
        return format_html('<a href="{}" target="_blank">Preview Report</a>', url)


    def admin_preview_report(self, request, report_id):
        report = get_object_or_404(ExperimentReport, id=report_id)
        
        # Access the related ReportEvaluation object via the related_name="evaluation"
        evaluation_obj = getattr(report, 'evaluation', None)
        
        # Extract the dictionaries from the JSONFields
        eval_scores = evaluation_obj.scores if evaluation_obj else {}
        eval_feedback = evaluation_obj.feedback if evaluation_obj else {}

        # Map the AI keys to the keys your report_detail.html template expects
        scores = {
            'objective': eval_scores.get('objective', 0),
            'theory': eval_scores.get('theory', 0),
            'apparatus_scope': eval_scores.get('apparatus_scope', 0),
            'procedure': eval_scores.get('procedure', 0),
            'results': eval_scores.get('results', 0),
            'data_analysis': eval_scores.get('data_analysis', 0),
            'discussion': eval_scores.get('discussion', 0),
            'conclusion': eval_scores.get('conclusion', 0),
            'references': eval_scores.get('references', 0),
        }

        feedback = {
            'objective': eval_feedback.get('objective', ""),
            'theory': eval_feedback.get('theory', ""),
            'apparatus_scope': eval_feedback.get('apparatus_scope', ""),
            'procedure': eval_feedback.get('procedure', ""),
            'results': eval_feedback.get('results', ""),
            'data_analysis': eval_feedback.get('data_analysis', ""),
            'discussion': eval_feedback.get('discussion', ""),
            'conclusion': eval_feedback.get('conclusion', ""),
            'references': eval_feedback.get('references', ""),
        }

        # Clean up any nested rubric objects for the radar chart
        for k, v in scores.items():
            if isinstance(v, dict):
                scores[k] = v.get("rubric", 0)

        context = {
            'report': report,
            'sections': [
                ('objective', 'Objective', 'bullseye'),
                ('theory', 'Theory', 'book'),
                ('apparatus_scope', 'Apparatus Used', 'tools'),
                ('procedure', 'Procedure', 'list-ol'),
                ('results', 'Results', 'chart-line'),
                ('data_analysis', 'Data & Error Analysis', 'search'),
                ('discussion', 'Discussion', 'comments'),
                ('conclusion', 'Conclusion', 'check-circle'),
                ('references', 'References', 'link'),
            ],
            'experimental_tables': [], 
            'observation_text': report.observation,
            'data_analysis_lines': report.data_analysis.split('\n') if report.data_analysis else [],
            'overall_score': evaluation_obj.total_score() if evaluation_obj else 0.0,
            'not_evaluated': False if evaluation_obj else True,
            'scores': scores,
            'feedback': feedback,
            'admin_preview': True,
        }

        return render(request, 'bank/report_detail.html', context)

    def evaluate_reports(self, request, queryset):
        import logging
        from QuestionBank.tasks import evaluate_reports_task

        logger = logging.getLogger(__name__)

        report_ids = list(queryset.values_list("id", flat=True))
        evaluator_id = request.user.id

        logger.info(f"Queuing evaluation for reports {report_ids}")

        evaluate_reports_task.delay(report_ids, evaluator_id)

        self.message_user(
            request,
            f"Evaluation task for {len(report_ids)} report(s) queued.",
            level=messages.INFO,
        )

    evaluate_reports.short_description = "Evaluate reports with Gemma 4"

    def has_module_permission(self, request):
            return False

    
@admin.register(YearReports)
class YearReportsAdmin(admin.ModelAdmin):
    list_display = ['year_label', 'open_year_link']

    def get_queryset(self, request):
        # Keeps the table empty as we are using a custom grid template
        return super().get_queryset(request).none() 

    def changelist_view(self, request, extra_context=None):
            extra_context = extra_context or {}
            year_data = []

            for year_id, year_name in YEAR_CHOICES:
                # Query the database for current submission totals
                report_count = ExperimentReport.objects.filter(
                    student__year_group=year_id, 
                    submitted_at__isnull=False
                ).count()

                # Identify the latest activity timestamp
                last_report = ExperimentReport.objects.filter(
                    student__year_group=year_id,
                    submitted_at__isnull=False
                ).order_by('-submitted_at').first()
                
                last_update = last_report.submitted_at.strftime("%b %d, %Y") if last_report else "No activity"

                year_data.append({
                    'id': year_id,
                    'name': year_name,
                    'count': report_count,
                    'last_update': last_update
                })

            extra_context['year_rows'] = year_data
            self.change_list_template = "admin/year_selection_list.html"
            return super().changelist_view(request, extra_context=extra_context)

    def year_label(self, obj):
        return "Academic Year Folder"

    def open_year_link(self, obj):
        return "Select a year group to begin audit"

@admin.register(ProgrammeReports)
class ProgrammeReportsAdmin(admin.ModelAdmin):
    list_display = ['name', 'code', 'open_folder_link']

    def changelist_view(self, request, extra_context=None):
        # Capture year from URL and store in session
        year_group = request.GET.get('year_group')
        if year_group:
            request.session['active_year_group'] = year_group
        return super().changelist_view(request, extra_context=extra_context)

    def open_folder_link(self, obj):
        # Pass both IDs to Level 3
        year_group = get_current_request().GET.get('year_group') or get_current_request().session.get('active_year_group')
        url = reverse("admin:virtuallab_programmeexperiment_changelist") + f"?programme_id={obj.id}"
        
        if year_group:
            url += f"&year_group={year_group}"
            
        return format_html(
            '<a class="button" href="{}" style="background-color: #417690; color: white;">'
            '<i class="fas fa-folder-open"></i> View Experiments</a>', 
            url
        )
    open_folder_link.short_description = "Navigation"
    def has_module_permission(self, request):
            # Keeps the URL active but hides it from the sidebar menu
            return False


@admin.register(ProgrammeExperiment)
class ProgrammeExperimentAdmin(admin.ModelAdmin):
    list_display = ['title', 'code', 'report_count_for_programme', 'view_student_reports']

    def changelist_view(self, request, extra_context=None):
        prog_id = request.GET.get('programme_id')
        year_group = request.GET.get('year_group')
        
        if prog_id: request.session['active_programme_id'] = prog_id
        if year_group: request.session['active_year_group'] = year_group
        
        # Dynamic Title for robustness
        active_prog = request.GET.get('programme_id') or request.session.get('active_programme_id')
        active_year = request.GET.get('year_group') or request.session.get('active_year_group')
        
        if active_prog:
            from .models import Programme
            prog = Programme.objects.filter(id=active_prog).first()
            year_label = dict(YEAR_CHOICES).get(active_year, "Unknown Year") if active_year else ""
            if prog:
                extra_context = extra_context or {}
                extra_context['title'] = f"Reports: {prog.name} - {year_label}"
        
        return super().changelist_view(request, extra_context=extra_context)

    def report_count_for_programme(self, obj):
        request = get_current_request()
        if not request: return "—"
        
        prog_id = request.GET.get('programme_id') or request.session.get('active_programme_id')
        year_group = request.GET.get('year_group') or request.session.get('active_year_group')

        if not prog_id:
            return format_html('<span style="color:orange;">Select Programme First</span>')
        
        filters = Q(experiment=obj, student__programme_id=prog_id, submitted_at__isnull=False)
        if year_group:
            filters &= Q(student__year_group=year_group) # Filter by the specific Year
        
        return ExperimentReport.objects.filter(filters).count()
    report_count_for_programme.short_description = "Submissions"

    def view_student_reports(self, obj):
        request = get_current_request()
        if not request: return "-"
        
        prog_id = request.GET.get('programme_id') or request.session.get('active_programme_id')
        year_group = request.GET.get('year_group') or request.session.get('active_year_group')

        if not prog_id: return "-"
        
        url = (
            reverse("admin:virtuallab_experimentreport_changelist")
            + f"?student__programme__id__exact={prog_id}&experiment__id__exact={obj.id}"
        )
        if year_group:
            url += f"&student__year_group__exact={year_group}" # Carry the Year Group to the final list
            
        return format_html(
            '<a class="button" href="{}" style="background-color: #79aec8; color: white;">'
            '<i class="fas fa-eye"></i> View Reports</a>', 
            url
        )

    def has_module_permission(self, request):
        # Keeps the URL active but hides it from the sidebar menu
        return False


class ExperimentAnswerInline(admin.TabularInline):
    model = ExperimentAnswer
    extra = 0
    readonly_fields = ('question', 'answer_text')


class ReportEvaluationInline(admin.StackedInline):
    model = ReportEvaluation
    can_delete = False
    extra = 0
    readonly_fields = ('evaluator', 'evaluated_at', 'display_scores', 'total_score')

    def display_scores(self, obj):
        if not obj or not obj.scores:
            return "No scores available"
        html = "<ul style='list-style:none; padding-left:0; margin:0;'>"
        for section, score in obj.scores.items():
            feedback = obj.feedback.get(section, "No feedback") if obj.feedback else "No feedback"
            html += f"<li><strong>{section.replace('_', ' ').title()}</strong>: {score}/10<br><em>{feedback}</em></li>"
        html += "</ul>"
        return mark_safe(html)
    display_scores.short_description = "Evaluation Scores & Feedback"

    def total_score(self, obj):
        if not obj or not obj.scores:
            return "-"
        total = 0
        for v in obj.scores.values(): 
            if isinstance(v, dict):
                try:
                    total += int(v.get("rubric", 0))
                except (TypeError, ValueError):
                    pass
        return total
    total_score.short_description = "Total Score"